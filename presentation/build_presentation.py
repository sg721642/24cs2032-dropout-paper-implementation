import os
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

PROJECT_ROOT = Path(__file__).resolve().parent.parent

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Colors
    BG_COLOR = RGBColor(255, 255, 255)
    TEXT_MAIN = RGBColor(15, 23, 42)      # Slate 900
    TEXT_MUTED = RGBColor(71, 85, 105)    # Slate 600
    ACCENT_NAVY = RGBColor(30, 58, 138)   # Blue 900
    ACCENT_BLUE = RGBColor(37, 99, 235)   # Blue 600
    BOX_BG = RGBColor(248, 250, 252)      # Slate 50
    BOX_BORDER = RGBColor(226, 232, 240)  # Slate 200
    CARD_BG = RGBColor(241, 245, 249)     # Slate 100
    GREEN_TEXT = RGBColor(22, 101, 52)    # Green 800
    RED_TEXT = RGBColor(153, 27, 27)      # Red 800

    def add_header(slide, title_text, category_text="24CS2032 DEEP LEARNING — PAPER IMPLEMENTATION STUDY"):
        # Category tracker
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        tf_cat.margin_left = tf_cat.margin_top = tf_cat.margin_right = tf_cat.margin_bottom = 0
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = ACCENT_BLUE

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.6))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_MAIN

        # Divider line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.4), Inches(11.733), Inches(0.02))
        line.fill.solid()
        line.fill.fore_color.rgb = BOX_BORDER
        line.line.color.rgb = BOX_BORDER

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    slide1 = prs.slides.add_slide(prs.slide_layouts[6])
    
    # Title box
    tbox = slide1.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(11.333), Inches(2.2))
    tf1 = tbox.text_frame
    tf1.word_wrap = True
    
    p0 = tf1.paragraphs[0]
    p0.text = "Manual Dropout in Neural Networks"
    p0.font.size = Pt(36)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_NAVY
    p0.space_after = Pt(12)
    
    p1 = tf1.add_paragraph()
    p1.text = "A Reduced-Scale Reproduction and Ablation Study"
    p1.font.size = Pt(20)
    p1.font.color.rgb = TEXT_MUTED
    p1.space_after = Pt(8)
    
    p2 = tf1.add_paragraph()
    p2.text = "Based on Srivastava et al., Journal of Machine Learning Research (2014)"
    p2.font.size = Pt(13)
    p2.font.italic = True
    p2.font.color.rgb = ACCENT_BLUE
    
    # Metadata Card
    card = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(4.3), Inches(11.333), Inches(2.2))
    card.fill.solid()
    card.fill.fore_color.rgb = BOX_BG
    card.line.color.rgb = BOX_BORDER
    card.line.width = Pt(1)
    
    ctf = card.text_frame
    ctf.word_wrap = True
    ctf.margin_left = Inches(0.4)
    ctf.margin_top = Inches(0.3)
    
    cp0 = ctf.paragraphs[0]
    cp0.text = "Student Presenter:"
    cp0.font.size = Pt(12)
    cp0.font.bold = True
    cp0.font.color.rgb = ACCENT_BLUE
    cp0.space_after = Pt(4)
    
    cp1 = ctf.add_paragraph()
    cp1.text = "Satyam Gupta  |  Roll No. 24CS2032"
    cp1.font.size = Pt(18)
    cp1.font.bold = True
    cp1.font.color.rgb = TEXT_MAIN
    cp1.space_after = Pt(4)
    
    cp2 = ctf.add_paragraph()
    cp2.text = "Integrated Dual Degree in Computer Science & Engineering and Artificial Intelligence"
    cp2.font.size = Pt(13)
    cp2.font.color.rgb = TEXT_MUTED
    cp2.space_after = Pt(2)
    
    cp3 = ctf.add_paragraph()
    cp3.text = "Rajiv Gandhi Institute of Petroleum Technology (RGIPT)"
    cp3.font.size = Pt(13)
    cp3.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 2: WHY DROPOUT?
    # =========================================================================
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide2, "Why Dropout? Addressing Co-adaptation & Overfitting")
    
    # Left Column: Concept & Problem
    lbox = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.2))
    lbox.fill.solid()
    lbox.fill.fore_color.rgb = BOX_BG
    lbox.line.color.rgb = BOX_BORDER
    ltf = lbox.text_frame
    ltf.word_wrap = True
    ltf.margin_left = ltf.margin_top = ltf.margin_right = ltf.margin_bottom = Inches(0.35)
    
    lp0 = ltf.paragraphs[0]
    lp0.text = "The Overfitting & Co-adaptation Problem"
    lp0.font.size = Pt(16)
    lp0.font.bold = True
    lp0.font.color.rgb = ACCENT_NAVY
    lp0.space_after = Pt(12)
    
    lp1 = ltf.add_paragraph()
    lp1.text = "• Deep networks have enormous expressive capacity. When trained on finite sample sizes, training error easily approaches zero while generalization error remains high."
    lp1.font.size = Pt(13)
    lp1.font.color.rgb = TEXT_MAIN
    lp1.space_after = Pt(10)
    
    lp2 = ltf.add_paragraph()
    lp2.text = "• Feature Co-adaptation: Neurons learn to compensate for the idiosyncratic errors of other specific neurons rather than forming robust, standalone feature representations."
    lp2.font.size = Pt(13)
    lp2.font.color.rgb = TEXT_MAIN
    lp2.space_after = Pt(10)
    
    lp3 = ltf.add_paragraph()
    lp3.text = "• Ensemble Limitation: Training and maintaining an ensemble of hundreds of large separate networks is computationally prohibitive."
    lp3.font.size = Pt(13)
    lp3.font.color.rgb = TEXT_MUTED
    lp3.space_after = Pt(10)
    
    lp4 = ltf.add_paragraph()
    lp4.text = "• Regularization Nuance: Dropout regularizes capacity, but excessive drop rates starve feature flow and cause underfitting."
    lp4.font.size = Pt(13)
    lp4.font.bold = True
    lp4.font.color.rgb = ACCENT_BLUE

    # Right Column: The 3-Step Dropout Mechanism
    rbox = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.8), Inches(1.7), Inches(5.733), Inches(5.2))
    rbox.fill.solid()
    rbox.fill.fore_color.rgb = BOX_BG
    rbox.line.color.rgb = BOX_BORDER
    rtf = rbox.text_frame
    rtf.word_wrap = True
    rtf.margin_left = rtf.margin_top = rtf.margin_right = rtf.margin_bottom = Inches(0.35)
    
    rp0 = rtf.paragraphs[0]
    rp0.text = "The 3-Phase Regularization Solution"
    rp0.font.size = Pt(16)
    rp0.font.bold = True
    rp0.font.color.rgb = ACCENT_NAVY
    rp0.space_after = Pt(14)
    
    steps = [
        ("1. Standard Forward Pass", "Hidden activations y = ReLU(W x + b) are computed across layers."),
        ("2. Stochastic Training Mask (p = drop rate, q = 1 - d)", "During training, units are dropped with probability d (retained with keep probability q). Forces each neuron to learn useful features alongside random sub-networks."),
        ("3. Deterministic Inference (Implicit Ensemble)", "At test time, the full network runs deterministically without dropping units, acting as an implicit geometric mean ensemble.")
    ]
    
    for title, desc in steps:
        p_t = rtf.add_paragraph()
        p_t.text = title
        p_t.font.size = Pt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = ACCENT_BLUE
        p_t.space_after = Pt(3)
        
        p_d = rtf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = TEXT_MAIN
        p_d.space_after = Pt(12)

    # =========================================================================
    # SLIDE 3: MANUAL DROPOUT IMPLEMENTATION
    # =========================================================================
    slide3 = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide3, "Manual Dropout: Mathematical Formulation & PyTorch Implementation")
    
    # Left Box: Mathematics & Expectation
    m_box = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.2))
    m_box.fill.solid()
    m_box.fill.fore_color.rgb = BOX_BG
    m_box.line.color.rgb = BOX_BORDER
    mtf = m_box.text_frame
    mtf.word_wrap = True
    mtf.margin_left = mtf.margin_top = mtf.margin_right = mtf.margin_bottom = Inches(0.35)
    
    mp0 = mtf.paragraphs[0]
    mp0.text = "Mathematical Formulation"
    mp0.font.size = Pt(16)
    mp0.font.bold = True
    mp0.font.color.rgb = ACCENT_NAVY
    mp0.space_after = Pt(8)
    
    eqs = [
        "Drop probability: d in [0, 1)   |   Keep probability: q = 1 - d",
        "Bernoulli mask sampling:   m_i ~ Bernoulli(q)",
        "Training pass (Inverted):   h̃_i = (m_i * h_i) / q",
        "Evaluation pass:           h̃_i = h_i   (Identity)"
    ]
    for eq in eqs:
        p = mtf.add_paragraph()
        p.text = eq
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = TEXT_MAIN
        p.space_after = Pt(6)
        
    p_exp_title = mtf.add_paragraph()
    p_exp_title.text = "Expectation Preservation Proof:"
    p_exp_title.font.size = Pt(13)
    p_exp_title.font.bold = True
    p_exp_title.font.color.rgb = ACCENT_BLUE
    p_exp_title.space_before = Pt(8)
    p_exp_title.space_after = Pt(4)
    
    p_proof = mtf.add_paragraph()
    p_proof.text = "E[h̃_i | h_i] = E[(m_i * h_i) / q | h_i] = (h_i / q) * E[m_i] = (h_i / q) * q = h_i"
    p_proof.font.size = Pt(12)
    p_proof.font.bold = True
    p_proof.font.color.rgb = GREEN_TEXT
    p_proof.space_after = Pt(8)
    
    p_exp_desc = mtf.add_paragraph()
    p_exp_desc.text = "Dividing by q keeps the expected activation magnitude unchanged during training, ensuring evaluation requires zero post-hoc weight scaling."
    p_exp_desc.font.size = Pt(12)
    p_exp_desc.font.color.rgb = TEXT_MUTED

    # Right Box: Real Code Excerpt from run.py
    c_box = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.8), Inches(1.7), Inches(5.733), Inches(5.2))
    c_box.fill.solid()
    c_box.fill.fore_color.rgb = RGBColor(15, 23, 42) # Dark editor background
    c_box.line.color.rgb = RGBColor(51, 65, 85)
    ctf2 = c_box.text_frame
    ctf2.word_wrap = True
    ctf2.margin_left = ctf2.margin_top = ctf2.margin_right = ctf2.margin_bottom = Inches(0.35)
    
    cp_title = ctf2.paragraphs[0]
    cp_title.text = "# run.py: ManualDropout Implementation"
    cp_title.font.size = Pt(14)
    cp_title.font.bold = True
    cp_title.font.color.rgb = RGBColor(148, 163, 184)
    cp_title.space_after = Pt(10)
    
    code_lines = [
        "class ManualDropout(nn.Module):",
        "    def __init__(self, drop_rate=0.5, inverted=True):",
        "        super().__init__()",
        "        self.drop_rate = float(drop_rate)",
        "        self.inverted = inverted",
        "",
        "    def forward(self, x):",
        "        keep = 1.0 - self.drop_rate",
        "        if not self.training:",
        "            return x if self.inverted else x * keep",
        "        if self.drop_rate == 0:",
        "            return x",
        "        # Bernoulli mask sampled on same device/dtype",
        "        mask = (torch.rand_like(x) < keep).to(x.dtype)",
        "        if self.inverted:",
        "            return x * mask / keep",
        "        return x * mask"
    ]
    for cl in code_lines:
        p = ctf2.add_paragraph()
        p.text = cl
        p.font.size = Pt(11)
        p.font.name = "Courier New"
        p.font.color.rgb = RGBColor(241, 245, 249)
        p.space_after = Pt(1)
        
    p_notes = ctf2.add_paragraph()
    p_notes.text = "\n[RULES] • No nn.Dropout  • Inverted scaling  • Hidden layers only"
    p_notes.font.size = Pt(10)
    p_notes.font.bold = True
    p_notes.font.color.rgb = RGBColor(56, 189, 248) # Sky blue

    # =========================================================================
    # SLIDE 4: CONTROLLED EXPERIMENTAL SETUP
    # =========================================================================
    slide4 = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide4, "Controlled Experimental Protocol (dropout-24cs2032-v1)")
    
    # Left Card: Architecture Specification
    arch_card = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.2))
    arch_card.fill.solid()
    arch_card.fill.fore_color.rgb = BOX_BG
    arch_card.line.color.rgb = BOX_BORDER
    atf = arch_card.text_frame
    atf.word_wrap = True
    atf.margin_left = atf.margin_top = atf.margin_right = atf.margin_bottom = Inches(0.35)
    
    ap0 = atf.paragraphs[0]
    ap0.text = "Architecture & Model Specification"
    ap0.font.size = Pt(16)
    ap0.font.bold = True
    ap0.font.color.rgb = ACCENT_NAVY
    ap0.space_after = Pt(12)
    
    specs = [
        ("Network Topology:", "784 -> 256 -> 256 -> 256 -> 10 Multilayer Perceptron"),
        ("Hidden Activation:", "Rectified Linear Unit (ReLU)"),
        ("Dropout Placement:", "Immediately after hidden ReLUs (Layers 1, 2, 3)"),
        ("Input/Output Dropout:", "NONE (Raw pixels & logits untouched)"),
        ("Parameter Count:", "Exactly 335,114 trainable weights & biases"),
        ("Hardware Backend:", "Apple Silicon CPU (Native arm64)"),
        ("Compute Benchmark:", "CPU: 0.785s / 2 eps  vs  MPS: 4.398s / 2 eps (>5.5x faster)")
    ]
    for lbl, val in specs:
        p = atf.add_paragraph()
        p.text = f"{lbl} {val}"
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MAIN
        p.space_after = Pt(7)

    # Right Card: Scientific Protocol Parameters
    proto_card = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.8), Inches(1.7), Inches(5.733), Inches(5.2))
    proto_card.fill.solid()
    proto_card.fill.fore_color.rgb = BOX_BG
    proto_card.line.color.rgb = BOX_BORDER
    ptf = proto_card.text_frame
    ptf.word_wrap = True
    ptf.margin_left = ptf.margin_top = ptf.margin_right = ptf.margin_bottom = Inches(0.35)
    
    pp0 = ptf.paragraphs[0]
    pp0.text = "Experimental Protocol & Sweep Design"
    pp0.font.size = Pt(16)
    pp0.font.bold = True
    pp0.font.color.rgb = ACCENT_NAVY
    pp0.space_after = Pt(12)
    
    pspecs = [
        ("Datasets:", "MNIST (primary) & Fashion-MNIST (extension)"),
        ("Data Splits:", "10,000 train / 5,000 val / 10,000 test (Split seed: 2026)"),
        ("Drop Rates Evaluated:", "8 uniform rates: d in {0.0, 0.1, 0.2, 0.3, 0.4, 0.5, 0.6, 0.7}"),
        ("Random Seeds:", "3 seeds: 42, 43, 44 (Bitwise-identical initialization per seed)"),
        ("Training Horizon:", "20 fixed epochs, batch size 128"),
        ("Optimizer:", "Adam (lr = 0.001, beta1=0.9, beta2=0.999, weight decay = 0)"),
        ("Model Selection Rule:", "Lowest mean validation error across seeds at epoch 20"),
        ("Evaluation Discipline:", "Official test set strictly quarantined (descriptive only)")
    ]
    for lbl, val in pspecs:
        p = ptf.add_paragraph()
        p.text = f"{lbl} {val}"
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MAIN
        p.space_after = Pt(5)
        
    p_vol = ptf.add_paragraph()
    p_vol.text = "\n[EXPERIMENT VOLUME] 24 runs & 480 epoch rows per dataset (48 runs total)"
    p_vol.font.size = Pt(11)
    p_vol.font.bold = True
    p_vol.font.color.rgb = ACCENT_BLUE

    # =========================================================================
    # SLIDE 5: MNIST RESULTS
    # =========================================================================
    slide5 = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide5, "MNIST: Dropout Reduced Test Error at Moderate Rates")
    
    # Left: Plot image
    mnist_plot_path = PROJECT_ROOT / "results_mnist" / "mnist_keep_probability.png"
    if mnist_plot_path.exists():
        slide5.shapes.add_picture(str(mnist_plot_path), Inches(0.8), Inches(1.7), Inches(5.8), Inches(5.2))
        
    # Right: Quantitative Summary Card
    res_card = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.9), Inches(1.7), Inches(5.633), Inches(5.2))
    res_card.fill.solid()
    res_card.fill.fore_color.rgb = BOX_BG
    res_card.line.color.rgb = BOX_BORDER
    rtf5 = res_card.text_frame
    rtf5.word_wrap = True
    rtf5.margin_left = rtf5.margin_top = rtf5.margin_right = rtf5.margin_bottom = Inches(0.3)
    
    rp0 = rtf5.paragraphs[0]
    rp0.text = "Measured Results on MNIST (Fixed Epoch 20)"
    rp0.font.size = Pt(15)
    rp0.font.bold = True
    rp0.font.color.rgb = ACCENT_NAVY
    rp0.space_after = Pt(8)
    
    res_lines = [
        ("No Dropout Baseline (d = 0.0):", "Train: 0.203 ± 0.200%  |  Val: 4.260 ± 0.533%\nTest Error: 4.320 ± 0.295%  |  Gap: 4.117 ± 0.101 pp"),
        ("Validation-Selected (d* = 0.4):", "Train: 0.323 ± 0.047%  |  Val: 3.747 ± 0.130%\nTest Error: 3.843 ± 0.282%  (Accuracy: 96.157%)\nGeneralization Gap: 3.520 ± 0.295 pp"),
        ("Net Gain over Baseline:", "+0.477 pp lower test error  (11.03% relative reduction)"),
        ("Paper Reference Rate (d = 0.5):", "Test Error: 3.883 ± 0.087%  (Gap: 3.183 pp)"),
        ("Model Selection Nuance:", "Lowest numerical test error occurred at d=0.3 (3.820 ± 0.118%). The observed mean test-error difference between d=0.3 and d=0.4 is only 0.023 percentage points. Test performance was not used for model selection."),
        ("Negative Finding:", "Higher drop rates degraded performance below baseline: d=0.6 (4.380%) and d=0.7 (5.110%).")
    ]
    
    for lbl, desc in res_lines:
        p_l = rtf5.add_paragraph()
        p_l.text = lbl
        p_l.font.size = Pt(12)
        p_l.font.bold = True
        p_l.font.color.rgb = ACCENT_BLUE if "Net Gain" not in lbl else GREEN_TEXT
        p_l.space_after = Pt(1)
        
        p_d = rtf5.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = TEXT_MAIN
        p_d.space_after = Pt(6)

    # =========================================================================
    # SLIDE 6: PAPER COMPARISON + FASHION EXTENSION
    # =========================================================================
    slide6 = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide6, "Historical Paper Comparison & Fashion-MNIST Extension")
    
    # Left Card: Paper Comparison
    comp_card = slide6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.2))
    comp_card.fill.solid()
    comp_card.fill.fore_color.rgb = BOX_BG
    comp_card.line.color.rgb = BOX_BORDER
    ctf6 = comp_card.text_frame
    ctf6.word_wrap = True
    ctf6.margin_left = ctf6.margin_top = ctf6.margin_right = ctf6.margin_bottom = Inches(0.3)
    
    cp0 = ctf6.paragraphs[0]
    cp0.text = "Comparison with Srivastava et al. (2014)"
    cp0.font.size = Pt(15)
    cp0.font.bold = True
    cp0.font.color.rgb = ACCENT_NAVY
    cp0.space_after = Pt(8)
    
    comp_points = [
        ("Paper Table 2 Historical Score:", "Srivastava et al. report ~1.25% test error on MNIST using a 3x1024 ReLU architecture with dropout."),
        ("Our Measured Selected Score:", "3.843 ± 0.282% test error (d = 0.4, 3x256 MLP)."),
        ("NOT Directly Comparable:", "Our experiment is not an apples-to-apples comparison:\n• Architecture: 256 units vs. 1,024-8,192 units\n• Training set: 10,000 subset vs. full 60,000\n• Budget: 20 epochs vs. hundreds of epochs\n• Optimizer: Adam vs. SGD + momentum + max-norm"),
        ("Qualitative Trend:", "Our reduced-scale rate sweep shows a qualitatively similar non-monotonic trend: moderate dropout performs better than very high dropout under this fixed training budget.")
    ]
    for lbl, desc in comp_points:
        p_l = ctf6.add_paragraph()
        p_l.text = lbl
        p_l.font.size = Pt(12)
        p_l.font.bold = True
        p_l.font.color.rgb = ACCENT_BLUE if "NOT" not in lbl else RED_TEXT
        p_l.space_after = Pt(1)
        
        p_d = ctf6.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = TEXT_MAIN
        p_d.space_after = Pt(6)

    # Right Card: Fashion-MNIST Extension
    fash_card = slide6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.8), Inches(1.7), Inches(5.733), Inches(5.2))
    fash_card.fill.solid()
    fash_card.fill.fore_color.rgb = BOX_BG
    fash_card.line.color.rgb = BOX_BORDER
    ftf6 = fash_card.text_frame
    ftf6.word_wrap = True
    ftf6.margin_left = ftf6.margin_top = ftf6.margin_right = ftf6.margin_bottom = Inches(0.3)
    
    fp0 = ftf6.paragraphs[0]
    fp0.text = "Fashion-MNIST Extension (Same Protocol)"
    fp0.font.size = Pt(15)
    fp0.font.bold = True
    fp0.font.color.rgb = ACCENT_NAVY
    fp0.space_after = Pt(8)
    
    fash_points = [
        ("No Dropout Baseline (d = 0.0):", "Train: 7.130 ± 0.177%  |  Val: 13.127 ± 0.247%\nTest Error: 14.663 ± 0.175%  |  Gap: 7.533 pp"),
        ("Validation-Selected (d* = 0.2):", "Train: 8.583 ± 0.414%  |  Val: 13.087 ± 0.408%\nTest Error: 14.347 ± 0.253%  (Accuracy: 85.653%)\nNet Improvement: +0.317 pp (2.16% relative)"),
        ("Critical Negative Finding:", "Canonical d = 0.5 underperformed baseline:\nTest error rose to 15.133 ± 0.325% (+0.47 pp worse than baseline). At d = 0.7, error spiked to 17.453%."),
        ("Key Scientific Insight:", "Under the fixed 20-epoch budget, higher dropout rates produced higher train and test error on Fashion-MNIST, consistent with stronger underfitting.")
    ]
    for lbl, desc in fash_points:
        p_l = ftf6.add_paragraph()
        p_l.text = lbl
        p_l.font.size = Pt(12)
        p_l.font.bold = True
        p_l.font.color.rgb = GREEN_TEXT if "Selected" in lbl else (RED_TEXT if "Negative" in lbl else ACCENT_BLUE)
        p_l.space_after = Pt(1)
        
        p_d = ftf6.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = TEXT_MAIN
        p_d.space_after = Pt(6)

    # =========================================================================
    # SLIDE 7: FAILURE CASES + LIMITATIONS
    # =========================================================================
    slide7 = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide7, "Failure Case Inspection and Methodological Limitations")
    
    # Left Card: Failure Images & Callouts (width 7.5 inches ~60%)
    fail_card = slide7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.6), Inches(7.5), Inches(5.4))
    fail_card.fill.solid()
    fail_card.fill.fore_color.rgb = BOX_BG
    fail_card.line.color.rgb = BOX_BORDER
    
    # MNIST Header & Callout in a compact text frame
    tbox_m = slide7.shapes.add_textbox(Inches(0.95), Inches(1.70), Inches(7.2), Inches(0.55))
    tf_m = tbox_m.text_frame
    tf_m.word_wrap = True
    tf_m.margin_left = tf_m.margin_top = tf_m.margin_right = tf_m.margin_bottom = Inches(0)
    p_m_lbl = tf_m.paragraphs[0]
    p_m_lbl.text = "MNIST Failure Cases (Selected d = 0.4):"
    p_m_lbl.font.size = Pt(12)
    p_m_lbl.font.bold = True
    p_m_lbl.font.color.rgb = ACCENT_BLUE
    
    p_m_call = tf_m.add_paragraph()
    p_m_call.text = "Callouts: True 2 → Pred 3 (conf 0.52)   |   True 9 → Pred 8 (conf 0.91)"
    p_m_call.font.size = Pt(10)
    p_m_call.font.color.rgb = TEXT_MUTED
    
    # MNIST Cropped Image (Row 1)
    mnist_fail_path = PROJECT_ROOT / "results_mnist" / "mnist_failure_cases.png"
    if mnist_fail_path.exists():
        p_mnist = slide7.shapes.add_picture(str(mnist_fail_path), Inches(0.95), Inches(2.30), Inches(7.2), Inches(1.50))
        p_mnist.crop_top = 70.0 / 1440.0
        p_mnist.crop_bottom = 920.0 / 1440.0

    # Fashion Header & Callout
    tbox_f = slide7.shapes.add_textbox(Inches(0.95), Inches(3.95), Inches(7.2), Inches(0.55))
    tf_f = tbox_f.text_frame
    tf_f.word_wrap = True
    tf_f.margin_left = tf_f.margin_top = tf_f.margin_right = tf_f.margin_bottom = Inches(0)
    p_f_lbl = tf_f.paragraphs[0]
    p_f_lbl.text = "Fashion-MNIST Failure Cases (Selected d = 0.2):"
    p_f_lbl.font.size = Pt(12)
    p_f_lbl.font.bold = True
    p_f_lbl.font.color.rgb = ACCENT_BLUE
    
    p_f_call = tf_f.add_paragraph()
    p_f_call.text = "Callouts: Coat → Pullover (conf 0.71)   |   Ankle boot → Sneaker (conf 0.99)"
    p_f_call.font.size = Pt(10)
    p_f_call.font.color.rgb = TEXT_MUTED
    
    # Fashion Cropped Image (Row 1)
    fashion_fail_path = PROJECT_ROOT / "results_fashion" / "fashion_failure_cases.png"
    if fashion_fail_path.exists():
        p_fash = slide7.shapes.add_picture(str(fashion_fail_path), Inches(0.95), Inches(4.55), Inches(7.2), Inches(1.50))
        p_fash.crop_top = 70.0 / 1440.0
        p_fash.crop_bottom = 920.0 / 1440.0

    # Observational footnote
    tbox_note = slide7.shapes.add_textbox(Inches(0.95), Inches(6.20), Inches(7.2), Inches(0.65))
    tf_note = tbox_note.text_frame
    tf_note.word_wrap = True
    tf_note.margin_left = tf_note.margin_top = tf_note.margin_right = tf_note.margin_bottom = Inches(0)
    p_note = tf_note.paragraphs[0]
    p_note.text = "Note: Errors reflect visual stroke & silhouette ambiguities at 28×28 resolution without spatial convolutional inductive bias."
    p_note.font.size = Pt(9.5)
    p_note.font.italic = True
    p_note.font.color.rgb = TEXT_MUTED

    # Right Card: Methodological Limitations (width 4.033 inches ~35%)
    lim_card = slide7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.5), Inches(1.6), Inches(4.033), Inches(5.4))
    lim_card.fill.solid()
    lim_card.fill.fore_color.rgb = BOX_BG
    lim_card.line.color.rgb = BOX_BORDER
    ltf7 = lim_card.text_frame
    ltf7.word_wrap = True
    ltf7.margin_left = ltf7.margin_top = ltf7.margin_right = ltf7.margin_bottom = Inches(0.3)
    
    lp0 = ltf7.paragraphs[0]
    lp0.text = "Methodological Constraints & Scope"
    lp0.font.size = Pt(14)
    lp0.font.bold = True
    lp0.font.color.rgb = ACCENT_NAVY
    lp0.space_after = Pt(8)
    
    limits = [
        ("• 3 random seeds only:", "N=3 reflects empirical sample variance but lacks statistical power for formal hypothesis testing."),
        ("• 10k training subset:", "Subsampled training split magnifies regularizer sensitivity relative to the full 60,000 dataset."),
        ("• Fixed 20-epoch budget:", "Higher dropout may require more optimization; this experiment intentionally keeps the same 20-epoch budget for every rate."),
        ("• Fully connected MLP:", "Lacks the translational invariance of convolutional networks."),
        ("• No significance test:", "Sample SD (ddof=1) reported; descriptive comparisons only without p-values."),
        ("• Fashion-MNIST is an extension:", "Exploratory benchmark under the identical protocol (not present in 2014 paper).")
    ]
    for lbl, desc in limits:
        p_l = ltf7.add_paragraph()
        p_l.text = lbl
        p_l.font.size = Pt(11)
        p_l.font.bold = True
        p_l.font.color.rgb = ACCENT_BLUE
        p_l.space_after = Pt(1)
        
        p_d = ltf7.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(10)
        p_d.font.color.rgb = TEXT_MAIN
        p_d.space_after = Pt(5)

    # =========================================================================
    # SLIDE 8: CONCLUSION + DEMO
    # =========================================================================
    slide8 = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide8, "Conclusion, Live Model Demo & Resources")
    
    # Left Card: Core Takeaways
    conc_card = slide8.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.2))
    conc_card.fill.solid()
    conc_card.fill.fore_color.rgb = BOX_BG
    conc_card.line.color.rgb = BOX_BORDER
    ctf8 = conc_card.text_frame
    ctf8.word_wrap = True
    ctf8.margin_left = ctf8.margin_top = ctf8.margin_right = ctf8.margin_bottom = Inches(0.3)
    
    cp0 = ctf8.paragraphs[0]
    cp0.text = "Key Research Conclusions"
    cp0.font.size = Pt(15)
    cp0.font.bold = True
    cp0.font.color.rgb = ACCENT_NAVY
    cp0.space_after = Pt(10)
    
    concs = [
        "1. Complete Mathematical Implementation: Inverted dropout was built and verified from scratch without any built-in library dependencies.",
        "2. Regularization Efficacy: Moderate validation-selected dropout lowered mean test error relative to the no-dropout baseline in our fixed protocol:\n   • MNIST: 4.320% -> 3.843% test error (11.03% relative reduction)\n   • Fashion-MNIST: 14.663% -> 14.347% test error (2.16% relative reduction)",
        "3. Critical Negative Finding: Canonical d = 0.5 underperformed the no-dropout baseline on Fashion-MNIST, and rates d >= 0.6 caused underfitting on both benchmarks.",
        "4. Methodological Lesson: Dropout strength cannot be assumed universally; it must be calibrated systematically against validation data."
    ]
    for c in concs:
        p = ctf8.add_paragraph()
        p.text = c
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_MAIN
        p.space_after = Pt(8)

    # Right Card: Demo, Repository & References
    demo_card = slide8.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.8), Inches(1.7), Inches(5.733), Inches(5.2))
    demo_card.fill.solid()
    demo_card.fill.fore_color.rgb = BOX_BG
    demo_card.line.color.rgb = BOX_BORDER
    dtf8 = demo_card.text_frame
    dtf8.word_wrap = True
    dtf8.margin_left = dtf8.margin_top = dtf8.margin_right = dtf8.margin_bottom = Inches(0.3)
    
    dp0 = dtf8.paragraphs[0]
    dp0.text = "Saved-Model Demonstration & Project Assets"
    dp0.font.size = Pt(15)
    dp0.font.bold = True
    dp0.font.color.rgb = ACCENT_NAVY
    dp0.space_after = Pt(10)
    
    demo_info = [
        ("Interactive Demo (presentation/demo.py):", 
         "• Loads saved checkpoint: mnist_d0.4_s42.pt on CPU\n• First 10 MNIST test samples: 10/10 correct in the demo.\n  (Demo sample only — not the reported full-test accuracy.)\n• Formal selected mean test accuracy remains 96.157%."),
        ("Private Project Repository:", 
         "https://github.com/sg721642/24cs2032-dropout-paper-implementation\n(Contains complete source, reproducible notebook, and audited CSV data)"),
        ("Verified Academic Report:", 
         "report/24CS2032_Dropout_Report.pdf (5-page IEEE conference format)"),
        ("Primary Citation:", 
         "Srivastava, N., Hinton, G., Krizhevsky, A., Sutskever, I., & Salakhutdinov, R. 'Dropout: A Simple Way to Prevent Neural Networks from Overfitting.' JMLR 15.56 (2014): 1929-1958.")
    ]
    for lbl, desc in demo_info:
        p_l = dtf8.add_paragraph()
        p_l.text = lbl
        p_l.font.size = Pt(12)
        p_l.font.bold = True
        p_l.font.color.rgb = ACCENT_BLUE
        p_l.space_after = Pt(1)
        
        p_d = dtf8.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = TEXT_MAIN
        p_d.space_after = Pt(7)

    # Save presentation
    out_path = PROJECT_ROOT / "presentation" / "24CS2032_Dropout_Presentation.pptx"
    prs.save(out_path)
    print(f"[*] Presentation successfully generated and saved to: {out_path}")
    print(f"[*] Slide count: {len(prs.slides)}")

if __name__ == "__main__":
    create_presentation()
